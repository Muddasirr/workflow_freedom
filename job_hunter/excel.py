from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from job_hunter.models import Job

HEADERS = [
    "Match Score",
    "Title",
    "Company",
    "Role Fit",
    "Work Mode",
    "Location",
    "Location Fit",
    "Pakistan Friendly",
    "Matched Skills",
    "Seniority",
    "Experience Fit",
    "Salary",
    "Job Type",
    "Posted",
    "Source",
    "Apply URL",
    "HR / Recruiter Email",
    "Email Source",
    "Other Emails",
    "Company Domain",
    "Summary",
    "Notes",
    "Applied?",
    "Applied On",
    "Status",
    "My Notes",
]

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
YES_FILL = PatternFill("solid", fgColor="C6EFCE")
MAYBE_FILL = PatternFill("solid", fgColor="FFEB9C")
NO_FILL = PatternFill("solid", fgColor="FFC7CE")
ALT_FILL = PatternFill("solid", fgColor="F7F9FC")
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)


def _row(job: Job) -> list[object]:
    return [
        job.score,
        job.title,
        job.company,
        ", ".join(job.role_fit),
        job.work_mode,
        job.location or ", ".join(job.location_restrictions),
        job.location_fit,
        job.pakistan_friendly,
        ", ".join(job.matched_skills),
        job.seniority,
        job.experience_fit,
        job.salary,
        job.job_type,
        job.posted_at,
        job.source,
        job.url,
        job.hr_email,
        job.email_source,
        ", ".join(e for e in job.emails if e != job.hr_email),
        job.company_domain,
        job.excerpt,
        job.notes,
        "",
        "",
        "",
        "",
    ]


def _style_header(ws: Worksheet) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28


def _style_body(ws: Worksheet) -> None:
    widths = {
        "A": 12,
        "B": 42,
        "C": 26,
        "D": 28,
        "E": 12,
        "F": 28,
        "G": 24,
        "H": 18,
        "I": 36,
        "J": 18,
        "K": 14,
        "L": 14,
        "M": 18,
        "N": 42,
        "O": 32,
        "P": 28,
        "Q": 36,
        "R": 22,
        "S": 55,
        "T": 40,
        "U": 12,
        "V": 14,
        "W": 16,
        "X": 30,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    wrap_cols = {2, 6, 9, 14, 15, 17, 19, 20, 24}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(HEADERS)):
        for cell in row:
            cell.border = THIN
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=cell.column in wrap_cols,
            )
            if cell.row % 2 == 0:
                if not cell.fill or cell.fill.fgColor is None or cell.fill.fgColor.rgb in {"00000000", "0"}:
                    cell.fill = ALT_FILL
        friendly = ws.cell(row[0].row, 8).value
        fill = {"Yes": YES_FILL, "Maybe": MAYBE_FILL, "No": NO_FILL}.get(str(friendly))
        if fill:
            ws.cell(row[0].row, 8).fill = fill
        url = ws.cell(row[0].row, 16).value
        if url:
            url_cell = ws.cell(row[0].row, 16)
            url_cell.hyperlink = str(url)
            url_cell.font = Font(color="0563C1", underline="single")
        email = ws.cell(row[0].row, 17).value
        if email:
            email_cell = ws.cell(row[0].row, 17)
            email_cell.hyperlink = f"mailto:{email}"
            email_cell.font = Font(color="0563C1", underline="single")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"
    ws.sheet_view.showGridLines = False


def _write_sheet(ws: Worksheet, jobs: list[Job]) -> None:
    ws.append(HEADERS)
    for job in jobs:
        ws.append(_row(job))
    _style_header(ws)
    if jobs:
        _style_body(ws)
    else:
        ws.append(["No matching jobs in this view yet."])


def write_workbook(jobs: list[Job], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    ranked = sorted(jobs, key=lambda j: (-j.score, j.company.lower(), j.title.lower()))
    karachi = [j for j in ranked if "Karachi" in j.location_fit or (j.pakistan_friendly == "Yes" and j.work_mode != "remote")]
    pakistan = [j for j in ranked if j.pakistan_friendly == "Yes"]
    remote_ok = [j for j in ranked if j.work_mode == "remote" and j.pakistan_friendly in {"Yes", "Maybe"}]
    restricted = [j for j in ranked if j.work_mode == "remote" and j.pakistan_friendly == "No"]
    missing_email = [j for j in ranked if not j.hr_email]
    with_email = [j for j in ranked if j.hr_email]
    verified = [j for j in with_email if "guess" not in (j.email_source or "").lower()]

    wb = Workbook()
    overview = wb.active
    overview.title = "Overview"
    overview["A1"] = "Job hunter — AI / Product / Software / Frontend engineer"
    overview["A1"].font = Font(bold=True, size=16, color="1F3A5F")
    overview["A2"] = f"Generated {date.today().isoformat()}"
    stats = [
        ("Total matching roles", len(ranked)),
        ("Pakistan-friendly (Yes)", len(pakistan)),
        ("Karachi physical / local", len(karachi)),
        ("Remote you can likely do from PK", len(remote_ok)),
        ("Restricted remote (outside PK geo-limits)", len(restricted)),
        ("With HR / recruiting email", len(with_email)),
        ("Verified (not guessed)", len(verified)),
        ("Still missing email", len(missing_email)),
    ]
    overview.append([])
    overview.append(["Snapshot", "Count"])
    overview["A4"].font = HEADER_FONT
    overview["A4"].fill = HEADER_FILL
    overview["B4"].font = HEADER_FONT
    overview["B4"].fill = HEADER_FILL
    for label, value in stats:
        overview.append([label, value])
    overview.append([])
    overview.append(["How to use"])
    overview.append(["1. Start with 'Pakistan Friendly' and 'Karachi and Pakistan' — highest chance of a reply."])
    overview.append(["2. Open Apply URL, then email HR if an address is listed. Always verify guessed emails."])
    overview.append(["3. Mark Applied? / Status on any sheet. Filters are already enabled."])
    overview.append(["4. Re-run: python -m job_hunter   (optional: HUNTER_API_KEY in .env for more emails)."])
    overview.column_dimensions["A"].width = 52
    overview.column_dimensions["B"].width = 14

    sheets = [
        ("All Matches", ranked),
        ("Emails", with_email),
        ("Verified Emails", verified),
        ("Karachi and Pakistan", pakistan),
        ("Remote OK from PK", remote_ok),
        ("Restricted Remote", restricted),
        ("Missing Email", missing_email),
    ]
    for name, subset in sheets:
        ws = wb.create_sheet(name)
        _write_sheet(ws, subset)

    wb.save(path)
    return path


def write_csv(jobs: list[Job], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    ranked = sorted(jobs, key=lambda j: (-j.score, j.company.lower(), j.title.lower()))
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        for job in ranked:
            writer.writerow(_row(job))
    return path


def write_emails_csv(contacts, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "Company",
                "Region",
                "City",
                "HR / Recruiter Email",
                "Email Source",
                "Other Emails",
                "Domain",
                "Careers / Apply URL",
                "Sample Role",
            ]
        )
        seen_email: set[str] = set()
        for c in contacts:
            emails = list(dict.fromkeys(c.emails))
            if not emails:
                continue
            primary = emails[0]
            # Prefer careers/hr if present
            for e in emails:
                local = e.split("@")[0]
                if local in {"careers", "hr", "jobs", "recruiting", "talent", "people", "hiring"}:
                    primary = e
                    break
            if primary in seen_email:
                continue
            seen_email.add(primary)
            others = [e for e in emails if e != primary]
            writer.writerow(
                [
                    c.name,
                    c.region,
                    c.city,
                    primary,
                    c.email_source,
                    ", ".join(others),
                    c.domain,
                    c.apply_url or c.careers_url,
                    c.sample_role,
                ]
            )
    verified_path = path.with_name(path.stem.replace("emails", "emails_verified") + path.suffix)
    if verified_path == path:
        verified_path = path.with_name("emails_verified_" + path.name)
    with verified_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "Company",
                "Region",
                "City",
                "HR / Recruiter Email",
                "Email Source",
                "Other Emails",
                "Domain",
                "Careers / Apply URL",
                "Sample Role",
            ]
        )
        seen_v: set[str] = set()
        with path.open(encoding="utf-8-sig", newline="") as src:
            for row in csv.DictReader(src):
                if "guess" in (row.get("Email Source") or "").lower():
                    continue
                email = (row.get("HR / Recruiter Email") or "").strip().lower()
                if not email or email in seen_v:
                    continue
                seen_v.add(email)
                writer.writerow(
                    [
                        row.get("Company", ""),
                        row.get("Region", ""),
                        row.get("City", ""),
                        row.get("HR / Recruiter Email", ""),
                        row.get("Email Source", ""),
                        row.get("Other Emails", ""),
                        row.get("Domain", ""),
                        row.get("Careers / Apply URL", ""),
                        row.get("Sample Role", ""),
                    ]
                )
    return path
